"""Annual export uses real scoped routes and remains a read-only operation."""
from io import BytesIO
from datetime import datetime
import unittest
from openpyxl import load_workbook
from test_inspection_plans import InspectionPlanTests
from inspection_plan_excel import build_annual_workbook


class InspectionPlanExcelTests(unittest.TestCase):
    setUpClass=classmethod(InspectionPlanTests.setUpClass.__func__)
    tearDownClass=classmethod(InspectionPlanTests.tearDownClass.__func__)
    setUp=InspectionPlanTests.setUp
    tearDown=InspectionPlanTests.tearDown
    change=InspectionPlanTests.change
    save=InspectionPlanTests.save

    def export(self, extra=''):
        return self.client.get('/vehicle/inspection-plans/export.xlsx?scope=company&year=2026&month=2026-10'+extra)

    def workbook(self, response):
        self.assertEqual(response.status_code,200,response.get_data()[:200])
        self.assertIn('attachment',response.headers['Content-Disposition'])
        self.assertIn('no-store',response.headers['Cache-Control'])
        return load_workbook(BytesIO(response.data))

    def test_fiscal_grid_typed_dates_and_read_only(self):
        self.save([self.change(mon='2026-04'),self.change(mon='2027-03',kind='ad_hoc')])
        before=self.m.VehicleInspectionPlanEvent.query.count()
        wb=self.workbook(self.export())
        self.assertEqual(wb.sheetnames,['年間点検計画','点検明細'])
        ws=wb.worksheets[0]
        self.assertEqual([ws.cell(4,c).value.month for c in range(3,15)],[4,5,6,7,8,9,10,11,12,1,2,3])
        self.assertEqual(ws['N4'].value.year,2027)
        self.assertTrue(all(ws.cell(4,c).alignment.horizontal=='center' for c in range(3,15)))
        self.assertEqual(ws.auto_filter.ref,'A4:N5')
        self.assertEqual([(f.colId,f.hiddenButton,f.showButton) for f in ws.auto_filter.filterColumn],[(i,True,False) for i in range(2,14)])
        self.assertEqual(ws.row_dimensions[4].height,25)
        for sheet in wb:
            self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
            self.assertEqual(sheet.page_setup.fitToWidth,1)
            self.assertEqual(sheet.page_setup.fitToHeight,0)
            self.assertTrue(sheet.print_title_rows)
        self.assertIn('3か月点検',ws['C5'].value)
        self.assertIn('突発修理',ws['N5'].value)
        self.assertEqual(ws.freeze_panes,'C5')
        self.assertEqual(ws.page_setup.orientation,'landscape')
        self.assertIsInstance(wb.worksheets[1]['E2'].value,datetime)
        self.assertEqual(self.m.VehicleInspectionPlanEvent.query.count(),before)

    def test_search_kind_state_and_no_match(self):
        self.vehicle.chassis_number='ＥＸＰＯＲＴ １２３';self.m.db.session.commit()
        self.save([self.change(mon='2026-12',kind='ad_hoc')])
        wb=self.workbook(self.export('&q=export123&kind=ad_hoc&state=unreserved'))
        self.assertEqual(wb.worksheets[0].max_row,5)
        self.assertIn('突発修理',wb.worksheets[0]['K5'].value)
        wb=self.workbook(self.export('&q=missing'))
        self.assertEqual(wb.worksheets[0].max_row,4)
        self.assertEqual(wb.worksheets[1].max_row,1)

    def test_company_office_and_foreign_focus_boundaries(self):
        m=self.m;self.vehicle.office='安城'
        foreign=m.Vehicle(company_code='FOREIGN',chassis_number='SECRET',office='安城')
        m.db.session.add_all([foreign,m.Vehicle(company_code='TEST',chassis_number='OTHER',office='小山')]);m.db.session.commit()
        wb=self.workbook(self.export('&office=安城'))
        self.assertEqual(wb.worksheets[0].max_row,5)
        self.assertEqual(self.export('&vehicle_id='+str(foreign.id)).status_code,404)
        self.assertEqual(self.export('&office=foreign-office').status_code,400)
        self.assertNotIn('SECRET',str([c.value for row in wb.worksheets[0] for c in row]))

    def test_export_requires_login_and_valid_period(self):
        self.assertEqual(self.export('&kind=bogus').status_code,400)
        self.assertEqual(self.export('&state=bogus').status_code,400)
        self.assertEqual(self.client.get('/vehicle/inspection-plans/export.xlsx?year=9999').status_code,400)
        with self.client.session_transaction() as s:s.clear()
        response=self.export()
        self.assertIn(response.status_code,(302,401))
        self.assertNotEqual(response.mimetype,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def test_multiple_entries_candidates_and_formula_injection(self):
        data={'year':2026,'today':'2026-10-26','scope':'company','office':'','focus_vehicle_id':None,
              'vehicles':[{'id':1,'number':'=1+1','office':'+evil','chassis_number':'x'}], 'plans':[]}
        for i in range(5):
            data['plans'].append(dict(vehicle_record_id=1,kind='12month',status='candidate',display_state='candidate',planned_month='2026-11',expiry_date='2026-11-30',factory='=HYPERLINK("evil")',note='@evil'))
        wb=load_workbook(build_annual_workbook(data,{}))
        self.assertEqual(wb.worksheets[0]['J5'].value.count('12か月点検'),5)
        self.assertEqual(wb.worksheets[1].max_row,6)
        for ws,addr in [(wb.worksheets[0],'A5'),(wb.worksheets[1],'I2')]:
            self.assertEqual(ws[addr].data_type,'s')

    def test_export_link_is_visible_with_working_href_in_all_views_before_js(self):
        from html.parser import HTMLParser
        from urllib.parse import urlsplit, parse_qs
        class Links(HTMLParser):
            def __init__(self):
                super().__init__(); self.exports=[]
            def handle_starttag(self, tag, attrs):
                attrs=dict(attrs)
                if tag=='a' and 'data-plan-export' in attrs:self.exports.append(attrs)
        for view in ('month','calendar','year'):
            with self.subTest(view=view):
                response=self.client.get('/vehicle/inspection-plans',query_string={
                    'scope':'company','month':'2027-01','view':view,'kind':'ad_hoc',
                    'state':'candidate','q':'TEST'})
                self.assertEqual(response.status_code,200)
                parser=Links();parser.feed(response.get_data(as_text=True))
                self.assertEqual(len(parser.exports),1)
                link=parser.exports[0]
                self.assertNotIn('hidden',link)
                params=parse_qs(urlsplit(link['href']).query)
                self.assertEqual(params['view'],[view])
                self.assertEqual(params['q'],['TEST'])
                self.assertEqual(params['kind'],['ad_hoc'])
                self.assertEqual(params['state'],['candidate'])
                workbook=self.workbook(self.client.get(link['href']))
                self.assertIn('2026年度',workbook.worksheets[0]['A1'].value)
