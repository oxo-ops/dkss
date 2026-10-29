"""Q02 vehicle inspection workflow boundaries on real Flask endpoints.

Disposable SQLite only. Ordinary inspections are completed through the route;
fixed-template state is seeded to isolate manager/repair authorization.
Does not validate PostgreSQL concurrency, browser CSRF or production.
"""
from identity_fixture import identity_for, PersonSelectionClient
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch


class Q02VehicleWorkflowBoundaryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix='dkss-q02-vehicle-flow-')
        from dkss_app import create_app, models
        from dkss_app.extensions import db, limiter
        cls.db, cls.models = db, models
        env = {'SECRET_KEY': 'q02-vehicle-flow-test-only', 'RENDER': '',
               'DATABASE_URL': 'sqlite:///' + str(Path(cls.temp.name) / 'test.db'),
               'OPENAI_API_KEY': '', 'SMTP_HOST': '', 'S3_BUCKET_NAME': '',
               'VAPID_PRIVATE_KEY': '', 'VAPID_PUBLIC_KEY': '',
               'AWS_EC2_METADATA_DISABLED': 'true'}
        with patch.dict(os.environ, env):
            cls.app = create_app({'TESTING': True, 'WTF_CSRF_ENABLED': False})
        cls.app.test_client_class = PersonSelectionClient
        limiter.enabled = False

    @classmethod
    def tearDownClass(cls):
        with cls.app.app_context():
            cls.db.session.remove()
            cls.db.engine.dispose()
        cls.temp.cleanup()

    def setUp(self):
        self.ctx = self.app.app_context()
        self.ctx.push()
        self.db.drop_all()
        self.db.create_all()
        m = self.models
        for company in ('A', 'B'):
            self.db.session.add(m.Company(company_code=company, company_name=company, active=True))
            for username, role in [('admin', 'admin'), ('checker', 'user'),
                                   ('reviewer', 'user'), ('outsider', 'user'), ('support', 'itc')]:
                self.db.session.add(m.User(company_code=company, username=username, role=role,
                    password='unused', last_name='Test', first_name=username,
                    password_changed_at='2026-10-27 10:00:00'))
        for company in ('A', 'B'):
            self.db.session.add(m.Vehicle(company_code=company, chassis_number=company+'-TEST', deleted=False, office='test'))
        self.db.session.commit()
        self.client = self.app.test_client()
        self.login('admin')

    def tearDown(self):
        self.db.session.remove()
        self.ctx.pop()

    def login(self, username, company='A'):
        user = self.models.User.query.filter_by(company_code=company, username=username).one()
        with self.client.session_transaction() as session:
            session.clear()
            session.update(account_identity=identity_for(company, username, app=(self.client).application), company_code=company, username=username, role=user.role,
                name='Test ' + username, password_changed_at=user.password_changed_at)

    def request(self, method, url, **kwargs):
        response = self.client.open(url, method=method, **kwargs)
        self.addCleanup(response.close)
        return response

    def record(self, company='A', general=False, fixed=False, requested=None):
        m = self.models
        vehicle = m.Vehicle.query.filter_by(company_code=company).one()
        items = [{'item_type': 'check', 'category': '点検', 'content': '確認',
                  'input_type': 'text', 'answer_required': False},
                 {'item_type': 'approval', 'approval_label': '確認',
                  'approval_allow_general': general}]
        checklist = m.Checklist(company_code=company, name=company+'-INSPECTION',
            target='車両管理', frequency_unit='day', display_type='month', active=True,
            items_json=json.dumps(items, ensure_ascii=False))
        self.db.session.add(checklist)
        self.db.session.commit()
        if fixed:
            # Persisted snapshot markers are read by the production endpoints.
            items[0]['fixed_template_code'] = 'daily_inspection_truck_trailer'
            snapshot = {'id': checklist.id, 'name': checklist.name, 'items': items}
            row = m.VehicleChecklistResult(company_code=company, checklist_id=checklist.id,
                vehicle_record_id=vehicle.id, year='2026', month='10', day='23',
                checked_by='Test checker', checked_by_username='checker', checked_date='2026-10-23 10:00',
                status='承認待ち', approvals_json='[]', answers_json='[]',
                checklist_snapshot_json=json.dumps(snapshot, ensure_ascii=False),
                operation_judgment_json=json.dumps({'status':'未判定', 'version':0,
                    'requested_usernames':list(requested or []),
                    'defects':[{'defect_no':1, 'item_no':'0', 'content':'合成不具合', 'status':'対応待ち'}]}, ensure_ascii=False))
            self.db.session.add(row)
            self.db.session.commit()
        else:
            self.login('checker', company)
            response = self.request('POST', f'/vehicle/checklists/{checklist.id}/complete',
                data={'vehicle_record_id':vehicle.id, 'year':'2026', 'month':'10', 'day':'23'})
            self.assertEqual(response.status_code, 302)
            row = m.VehicleChecklistResult.query.filter_by(checklist_id=checklist.id).one()
            self.assertEqual(row.status, '承認待ち')
        self.login('admin')
        return row

    def url(self, row, suffix):
        return f'/vehicle/checklist-results/{row.id}' + suffix

    def snapshot(self):
        self.db.session.expire_all()
        return tuple(tuple(tuple(getattr(row, c.name) for c in model.__table__.columns)
                           for row in model.query.order_by(model.id).all())
                     for model in (self.models.VehicleChecklistResult, self.models.ChecklistEvent,
                                   self.models.Notification, self.models.AuditLog,
                                   self.models.VehicleMaintenanceRecord,
                                   self.models.VehicleMaintenanceRecordEvent))

    def judge_data(self, row, **overrides):
        self.db.session.refresh(row)
        return {'decision':'判定保留', 'reason':'確認中', 'authority_role':'整備管理者',
                'authority_confirmed':'1', 'judgment_version':str(json.loads(row.operation_judgment_json)['version']),
                'expected_answers':row.answers_json, 'expected_snapshot':row.checklist_snapshot_json,
                **overrides}

    def judge(self, row, **overrides):
        return self.request('POST', self.url(row, '/operation-judgment'), data=self.judge_data(row, **overrides))

    def repair_data(self, **overrides):
        return {'judgment_version':'0', 'repair_status':'再確認待ち', 'repair_note':'合成試験の整備',
                'repair_performed_date':'2026-10-23', **overrides}

    def repair(self, row, **overrides):
        return self.request('POST', self.url(row, '/defects/1/repair'), data=self.repair_data(**overrides))

    def test_foreign_result_excel_is_not_exported(self):
        row = self.record(company='B')
        before = self.snapshot()
        self.assertEqual(self.request('GET', self.url(row, '/excel')).status_code, 302)
        self.assertEqual(self.snapshot(), before)

    def test_foreign_result_cannot_be_approved_or_rejected(self):
        row = self.record(company='B')
        before = self.snapshot()
        for suffix in ('/approve/0', '/reject'):
            self.assertEqual(self.request('POST', self.url(row, suffix), data={'reject_reason':'bad'}).status_code, 302)
            self.assertEqual(self.snapshot(), before)

    def test_foreign_fixed_result_cannot_be_judged_repaired_or_assigned(self):
        row = self.record(company='B', fixed=True, requested=['admin'])
        before = self.snapshot()
        self.assertEqual(self.judge(row).status_code, 404)
        self.assertEqual(self.repair(row).status_code, 404)
        self.assertEqual(self.judge(row, action='assign_operation_manager',manager_username='admin').status_code,404)
        self.assertEqual(self.snapshot(), before)

    def test_general_user_cannot_approve_or_reject_admin_only_slot(self):
        row = self.record()
        self.login('reviewer')
        before = self.snapshot()
        for suffix in ('/approve/0','/reject'):
            self.assertEqual(self.request('POST',self.url(row,suffix),data={'reject_reason':'bad'}).status_code,302)
            self.assertEqual(self.snapshot(),before)

    def test_general_slot_accepts_user_and_ignores_forged_actor(self):
        row = self.record(general=True)
        self.login('reviewer')
        response = self.request('POST',self.url(row,'/approve/0'),
            data={'company_code':'B','approved_by_username':'admin','username':'admin'})
        self.assertEqual(response.status_code,302)
        self.db.session.refresh(row)
        self.assertEqual(row.status,'承認済み')
        self.assertEqual(row.approved_by_username,'reviewer')

    def test_invalid_index_and_replayed_approval_do_not_write(self):
        row = self.record()
        before = self.snapshot()
        self.request('POST',self.url(row,'/approve/99'))
        self.assertEqual(self.snapshot(),before)
        self.request('POST',self.url(row,'/approve/0'))
        self.db.session.refresh(row)
        self.assertEqual(row.status,'承認済み')
        before = self.snapshot()
        self.request('POST',self.url(row,'/approve/0'))
        self.assertEqual(self.snapshot(),before)

    def test_revoked_admin_cannot_approve_or_judge_using_stale_cookie(self):
        normal = self.record()
        fixed = self.record(fixed=True,requested=['admin'])
        self.models.User.query.filter_by(company_code='A',username='admin').one().role='user'
        self.db.session.commit()
        before = self.snapshot()
        self.request('POST',self.url(normal,'/approve/0'))
        self.assertEqual(self.judge(fixed).status_code,403)
        self.assertEqual(self.snapshot(),before)

    def test_itc_cannot_judge_or_register_repairs(self):
        row = self.record(fixed=True,requested=['support'])
        self.login('support')
        before = self.snapshot()
        self.assertEqual(self.judge(row).status_code,403)
        self.assertEqual(self.repair(row).status_code,403)
        self.assertEqual(self.snapshot(),before)

    def test_unselected_admin_cannot_judge(self):
        row = self.record(fixed=True,requested=['other-manager'])
        before = self.snapshot()
        self.assertEqual(self.judge(row).status_code,403)
        self.assertEqual(self.snapshot(),before)

    def test_selected_admin_can_judge_with_server_actor(self):
        row = self.record(fixed=True,requested=['admin'])
        self.assertEqual(self.judge(row,company_code='B',judged_by_username='reviewer').status_code,302)
        self.db.session.refresh(row)
        judgment=json.loads(row.operation_judgment_json)
        self.assertEqual(judgment['status'],'判定保留')
        self.assertEqual(judgment['judged_by_username'],'admin')
        event=self.models.ChecklistEvent.query.filter_by(result_id=row.id,event_type='運行判断').one()
        self.assertEqual((event.company_code,event.actor_username),('A','admin'))

    def test_stale_judgment_and_changed_inspection_evidence_are_rejected(self):
        row = self.record(fixed=True,requested=['admin'])
        before = self.snapshot()
        for data in ({'judgment_version':'999'},{'expected_answers':'[{"value":"forged"}]'},
                     {'expected_snapshot':'{}'}):
            with self.subTest(data=data):
                self.assertEqual(self.judge(row,**data).status_code,409)
                self.assertEqual(self.snapshot(),before)

    def test_judgment_requires_authority_confirmation(self):
        row = self.record(fixed=True,requested=['admin'])
        before = self.snapshot()
        self.assertEqual(self.judge(row,authority_confirmed='0').status_code,403)
        self.assertEqual(self.snapshot(),before)

    def test_foreign_and_nonadmin_notification_targets_are_rejected(self):
        self.db.session.add(self.models.User(company_code='B',username='foreign-only',role='admin',
            password='unused',last_name='Test',first_name='Foreign'))
        self.db.session.commit()
        row = self.record(fixed=True,requested=['admin'])
        before=self.snapshot()
        for recipient in ('foreign-only','reviewer'):
            self.assertEqual(self.judge(row,decision='運行不可',
                approval_notify_users_operation=recipient).status_code,400)
            self.assertEqual(self.snapshot(),before)

    def test_checker_can_assign_own_company_admin_only_once(self):
        row = self.record(fixed=True)
        self.login('checker')
        self.assertEqual(self.judge(row,action='assign_operation_manager',manager_username='admin').status_code,302)
        self.db.session.refresh(row)
        self.assertEqual(json.loads(row.operation_judgment_json)['requested_usernames'],['admin'])
        before=self.snapshot()
        self.assertEqual(self.judge(row,action='assign_operation_manager',manager_username='admin').status_code,409)
        self.assertEqual(self.snapshot(),before)

    def test_unrelated_user_cannot_assign_manager(self):
        row = self.record(fixed=True)
        self.login('outsider')
        before=self.snapshot()
        self.assertEqual(self.judge(row,action='assign_operation_manager',manager_username='admin').status_code,403)
        self.assertEqual(self.snapshot(),before)

    def test_checker_cannot_assign_nonadmin_or_foreign_only_manager(self):
        self.db.session.add(self.models.User(company_code='B',username='foreign-only',role='admin',
            password='unused',last_name='Test',first_name='Foreign'))
        self.db.session.commit()
        row = self.record(fixed=True)
        self.login('checker')
        before=self.snapshot()
        for manager in ('reviewer','foreign-only'):
            self.assertEqual(self.judge(row,action='assign_operation_manager',manager_username=manager).status_code,400)
            self.assertEqual(self.snapshot(),before)

    def test_repair_and_recheck_use_server_actor_and_preserve_other_company(self):
        row = self.record(fixed=True,requested=['admin'])
        foreign = self.record(company='B',fixed=True,requested=['admin'])
        foreign_before=foreign.operation_judgment_json
        self.login('reviewer')
        self.assertEqual(self.repair(row,company_code='B',performed_by_username='admin').status_code,302)
        self.db.session.refresh(row)
        defect=json.loads(row.operation_judgment_json)['defects'][0]
        self.assertEqual(defect['status'],'再確認待ち')
        self.assertEqual(defect['repair']['performed_by_username'],'reviewer')
        self.assertEqual(self.repair(row,judgment_version='1',recheck_result='異常なし').status_code,302)
        self.db.session.refresh(row)
        defect=json.loads(row.operation_judgment_json)['defects'][0]
        self.assertEqual(defect['status'],'解消')
        self.assertEqual(defect['recheck']['performed_by_username'],'reviewer')
        self.db.session.refresh(foreign)
        self.assertEqual(foreign.operation_judgment_json,foreign_before)
        self.assertTrue(all(r.company_code=='A' for r in self.models.VehicleMaintenanceRecord.query.all()))

    def test_stale_repair_and_unknown_defect_have_no_side_effects(self):
        row = self.record(fixed=True)
        before=self.snapshot()
        self.assertEqual(self.repair(row,judgment_version='999').status_code,409)
        self.assertEqual(self.request('POST',self.url(row,'/defects/999/repair'),data=self.repair_data()).status_code,404)
        self.assertEqual(self.snapshot(),before)

    def test_inconsistent_foreign_parent_blocks_fixed_workflow_writes(self):
        row = self.record(fixed=True,requested=['admin'])
        row.vehicle_record_id=self.models.Vehicle.query.filter_by(company_code='B').one().id
        self.db.session.commit()
        before=self.snapshot()
        self.assertEqual(self.judge(row).status_code,404)
        self.assertEqual(self.repair(row).status_code,404)
        self.assertEqual(self.snapshot(),before)

    def test_fixed_template_cannot_bypass_judgment_with_generic_approval(self):
        row = self.record(fixed=True,requested=['admin'])
        before=self.snapshot()
        self.assertEqual(self.request('POST',self.url(row,'/approve/0')).status_code,409)
        self.assertEqual(self.snapshot(),before)

    def test_operation_clearance_requires_checks_and_current_vehicle_evidence(self):
        row=self.record(fixed=True,requested=['admin'])
        before=self.snapshot()
        self.assertEqual(self.judge(row,decision='運行可').status_code,400)
        self.assertEqual(self.judge(row,decision='運行可',checks_confirmed='1',
            vehicle_operation_basis='stale').status_code,409)
        self.assertEqual(self.snapshot(),before)

    def test_selected_admin_can_record_clearance_with_current_evidence(self):
        row=self.record(fixed=True,requested=['admin'])
        state=json.loads(row.operation_judgment_json)
        state['defects']=[]
        row.operation_judgment_json=json.dumps(state,ensure_ascii=False)
        self.db.session.commit()
        from dkss_app.checklists import get_vehicle_operation_basis
        basis=get_vehicle_operation_basis('A',row.vehicle_record_id)
        self.assertEqual(self.judge(row,decision='運行可',checks_confirmed='1',
            vehicle_operation_basis=basis).status_code,302)
        self.db.session.refresh(row)
        judgment=json.loads(row.operation_judgment_json)
        self.assertEqual(judgment['status'],'運行可')
        self.assertTrue(judgment['checks_confirmed'])
        self.assertEqual(row.status,'承認済み')

    def test_validation_only_judgment_does_not_save_or_notify(self):
        row=self.record(fixed=True,requested=['admin'])
        before=self.snapshot()
        response=self.request('POST',self.url(row,'/operation-judgment'),
            data=self.judge_data(row),headers={'X-DKSS-Validation-Only':'1'})
        self.assertEqual(response.status_code,200)
        self.assertTrue(response.json['success'])
        self.assertEqual(self.snapshot(),before)

    def test_repair_replay_does_not_duplicate_maintenance_or_notifications(self):
        row=self.record(fixed=True)
        self.assertEqual(self.repair(row).status_code,302)
        before=self.snapshot()
        self.assertEqual(self.repair(row).status_code,409)
        self.assertEqual(self.snapshot(),before)

    def test_own_company_ordinary_result_remains_exportable(self):
        row=self.record()
        self.login('reviewer')
        response=self.request('GET',self.url(row,'/excel'))
        self.assertEqual(response.status_code,200)
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(response.data),read_only=True)
        try:
            self.assertTrue(book.sheetnames)
            self.assertTrue(any(sheet.max_row for sheet in book))
        finally:
            book.close()


if __name__ == '__main__':
    unittest.main(verbosity=2)
