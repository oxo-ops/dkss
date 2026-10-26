"""Excel representation of the same company-scoped annual inspection data."""
from datetime import date
from io import BytesIO
import unicodedata
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import FilterColumn

STATES = {'candidate':'予定未登録', 'unreserved':'未予約', 'reserved':'予約済み',
          'awaiting':'実施確認待ち', 'completed':'実施済み'}
KINDS = {'3month':'3か月点検', '12month':'12か月点検', 'ad_hoc':'突発点検'}
COLORS = {'candidate':'FFF3DF', 'unreserved':'F3F6FA', 'reserved':'E7F0FF',
          'awaiting':'FFF0DE', 'completed':'E5F4EB'}


def normalize(value):
    return ''.join(unicodedata.normalize('NFKC', str(value or '')).lower().split())


def annual_rows(data, args):
    """Match browser search/focus/state membership, including empty vehicle rows."""
    search = normalize(args.get('q'))
    focus = data.get('focus_vehicle_id')
    vehicles = [v for v in data['vehicles'] if (not focus or v['id'] == focus) and
                (not search or search in normalize(v['number']+' '+v['chassis_number']))]
    ids = {v['id'] for v in vehicles}
    rows = [p for p in data['plans'] if p['vehicle_record_id'] in ids and
            p['status'] != 'cancelled' and (not args.get('kind') or p['kind'] == args['kind']) and
            (not args.get('state') or p['display_state'] == args['state'])]
    if args.get('state'):
        visible = {p['vehicle_record_id'] for p in rows}
        vehicles = [v for v in vehicles if v['id'] in visible]
    return vehicles, rows


def rank(row, today):
    priority = 3 if row['status'] == 'completed' else 0 if row['display_state'] == 'awaiting' or (row.get('expiry_date') and row['expiry_date'] <= today) else 1 if not row.get('reservation_date') else 2
    return (priority, row.get('expiry_date') or row.get('reservation_date') or row['planned_month'], row['vehicle_record_id'])


def write_text(cell, value):
    # User-entered registration/factory/note must never become an Excel formula.
    cell.value = str(value or '')
    cell.data_type = 's'


def build_annual_workbook(data, args):
    vehicles, rows = annual_rows(data, args)
    rows = sorted(rows, key=lambda p: rank(p, data['today']))
    year = data['year']
    months = [f'{year+(i>=9):04d}-{(i+3)%12+1:02d}' for i in range(12)]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '年間点検計画'
    sheet.merge_cells('A1:N1')
    sheet['A1'] = f'{year}年度 点検計画（4月〜翌3月）'
    sheet['A1'].font = Font(name='Meiryo', size=16, bold=True, color='24476B')
    sheet.merge_cells('A2:N2')
    scope = '会社全体' if data['scope']=='company' else data['office'] or '営業所未設定'
    filters = [scope, KINDS.get(args.get('kind'), '点検種別すべて'), STATES.get(args.get('state'), '状態すべて')]
    if args.get('q'): filters.append('検索：'+args['q'])
    write_text(sheet['A2'], ' ／ '.join(filters))
    sheet.merge_cells('A3:N3')
    sheet['A3'] = f'{len(vehicles)}台 ／ 対象{len(rows)}件 ／ 出力日 {data["today"]}'
    sheet.append(['車両', '営業所', *[date.fromisoformat(m+'-01') for m in months]])
    for cell in sheet[4][2:]: cell.number_format = 'm"月"'
    for index, vehicle in enumerate(vehicles, 5):
        write_text(sheet.cell(index,1),vehicle['number'])
        write_text(sheet.cell(index,2),vehicle['office'])
        for col, mon in enumerate(months, 3):
            items = [p for p in rows if p['vehicle_record_id']==vehicle['id'] and
                     (p['planned_month']==mon or (p.get('performed_date') or '')[:7]==mon)]
            lines = []
            for p in items:
                lines.append(KINDS[p['kind']]+' '+STATES[p['display_state']])
                for field, label in [('reservation_date','予約'),('performed_date','実施'),('expiry_date','期限')]:
                    if p.get(field): lines.append(label+' '+p[field])
            cell=sheet.cell(index,col)
            write_text(cell,'\n'.join(lines))
            if items: cell.fill=PatternFill('solid',fgColor=COLORS[items[0]['display_state']])
        sheet.row_dimensions[index].height=max(46, max((len(str(sheet.cell(index,c).value or '').splitlines()) for c in range(3,15)),default=1)*15+8)
    details=workbook.create_sheet('点検明細')
    details.append(['車両','営業所','点検種別','状態','予定月','予約日','実施日','期限','実施先','備考'])
    by_id={v['id']:v for v in vehicles}
    for row_index, p in enumerate(rows, 2):
        v=by_id[p['vehicle_record_id']]
        values=[v['number'],v['office'],KINDS[p['kind']],STATES[p['display_state']],p['planned_month'],
                p.get('reservation_date'),p.get('performed_date'),p.get('expiry_date'),p.get('factory'),p.get('note')]
        for col,value in enumerate(values,1):
            cell=details.cell(row_index,col)
            if col in (5,6,7,8) and value:
                cell.value=date.fromisoformat(value+'-01' if col==5 else value)
                cell.number_format='yyyy-mm' if col==5 else 'yyyy-mm-dd'
            else: write_text(cell,value)
        details.row_dimensions[row_index].height=44
    line=Side(style='thin',color='D9E3EF')
    for ws,header in [(sheet,4),(details,1)]:
        ws.sheet_view.showGridLines=False
        for row in ws.iter_rows(min_row=header):
            for cell in row:
                cell.font=Font(name='Meiryo',size=10,color='243B53')
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                cell.border=Border(left=line,right=line,top=line,bottom=line)
        for cell in ws[header]:
            cell.fill=PatternFill('solid',fgColor='E8F0FC')
            cell.font=Font(name='Meiryo',size=10,bold=True,color='24476B')
        ws.freeze_panes='C5' if ws==sheet else 'C2'
        ws.auto_filter.ref=f'A{header}:{get_column_letter(ws.max_column)}{max(header,ws.max_row)}'
        ws.print_title_rows=f'1:{header}'
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape'
        ws.page_setup.paperSize=ws.PAPERSIZE_A3 if ws==sheet else ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth=1
        ws.page_setup.fitToHeight=0
        ws.print_options.horizontalCentered=True
        ws.print_area=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
        for col in range(1,ws.max_column+1):
            ws.column_dimensions[get_column_letter(col)].width=24 if col==1 else 18 if col==2 else 21
    # Keep the full filter/sort range so monthly entries stay with their vehicle.
    # Only vehicle/office need a dropdown; month headings must remain unobscured.
    for col_id in range(2,14):
        sheet.auto_filter.filterColumn.append(FilterColumn(colId=col_id,hiddenButton=True,showButton=False))
    sheet.row_dimensions[4].height=25
    for cell in sheet[4][2:]:
        cell.alignment=Alignment(horizontal='center',vertical='center')
    details.column_dimensions['I'].width=28
    details.column_dimensions['J'].width=45
    output=BytesIO();workbook.save(output);output.seek(0)
    return output
